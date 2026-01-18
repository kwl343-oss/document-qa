import streamlit as st
from openai import OpenAI
import json
from datetime import datetime

EXTRACT_SCHEMA = {
    "company": None,
    "stage": None,
    "sector": None,
    "raise_amount_usd": None,
    "arr_usd": None,
    "growth_rate_pct": None,
    "runway_months": None,
    "notes": None
}

st.set_page_config(
    page_title="VCaaS – Deal Dashboard",
    layout="wide",
    initial_sidebar_state="expanded"
)

# ----------------------------
# Premium Custom CSS
# ----------------------------
st.markdown("""
    <style>
    :root {
        --primary: #0052CC;
        --success: #28a745;
        --warning: #ffc107;
        --danger: #dc3545;
        --dark: #1f2937;
        --light: #f9fafb;
    }
    
    .main-title {
        font-size: 2.8em;
        font-weight: 700;
        background: linear-gradient(135deg, #0052CC 0%, #3385ff 100%);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        background-clip: text;
        margin-bottom: 0.5rem;
    }
    
    .deal-snapshot {
        background: linear-gradient(135deg, #ffffff 0%, #f8f9fa 100%);
        border: 1px solid #e0e0e0;
        border-radius: 12px;
        padding: 2rem;
        margin-bottom: 2rem;
        box-shadow: 0 4px 12px rgba(0,0,0,0.08);
    }
    
    .snapshot-header {
        font-size: 1.6em;
        font-weight: 700;
        color: #1f2937;
        margin-bottom: 0.5rem;
    }
    
    .snapshot-tagline {
        font-size: 0.95em;
        color: #6b7280;
        margin-bottom: 1.5rem;
    }
    
    .snapshot-row {
        display: grid;
        grid-template-columns: repeat(4, 1fr);
        gap: 1.5rem;
        margin-bottom: 1rem;
    }
    
    .snapshot-item {
        background: white;
        padding: 1rem;
        border-radius: 8px;
        border-left: 3px solid #0052CC;
    }
    
    .snapshot-label {
        font-size: 0.8em;
        color: #6b7280;
        text-transform: uppercase;
        letter-spacing: 0.5px;
        margin-bottom: 0.25rem;
        font-weight: 600;
    }
    
    .snapshot-value {
        font-size: 1.3em;
        font-weight: 700;
        color: #1f2937;
    }
    
    .recommendation-card {
        background: linear-gradient(135deg, #f0f7ff 0%, #ffffff 100%);
        border: 2px solid #0052CC;
        border-radius: 12px;
        padding: 2rem;
        margin-bottom: 2rem;
        text-align: center;
    }
    
    .recommendation-label {
        font-size: 0.9em;
        color: #6b7280;
        text-transform: uppercase;
        letter-spacing: 1px;
        margin-bottom: 0.5rem;
    }
    
    .recommendation-badge-invest {
        display: inline-block;
        background: #d4edda;
        color: #155724;
        padding: 1rem 2rem;
        border-radius: 8px;
        font-size: 1.4em;
        font-weight: 700;
        margin: 0.5rem 0.25rem;
    }
    
    .recommendation-badge-watchlist {
        display: inline-block;
        background: #fff3cd;
        color: #856404;
        padding: 1rem 2rem;
        border-radius: 8px;
        font-size: 1.4em;
        font-weight: 700;
        margin: 0.5rem 0.25rem;
    }
    
    .recommendation-badge-pass {
        display: inline-block;
        background: #f8d7da;
        color: #721c24;
        padding: 1rem 2rem;
        border-radius: 8px;
        font-size: 1.4em;
        font-weight: 700;
        margin: 0.5rem 0.25rem;
    }
    
    .confidence-low { color: #dc3545; font-weight: 700; }
    .confidence-med { color: #ffc107; font-weight: 700; }
    .confidence-high { color: #28a745; font-weight: 700; }
    
    .thesis-fit-score {
        background: linear-gradient(135deg, #f0f7ff 0%, #ffffff 100%);
        border: 1px solid #d4e6ff;
        border-radius: 12px;
        padding: 1.5rem;
        margin-bottom: 1.5rem;
    }
    
    .fit-gauge {
        font-size: 3em;
        font-weight: 700;
        text-align: center;
        margin: 1rem 0;
    }
    
    .fit-score-good { color: #28a745; }
    .fit-score-med { color: #ffc107; }
    .fit-score-bad { color: #dc3545; }
    
    .fit-item {
        display: flex;
        align-items: center;
        margin: 0.8rem 0;
        font-size: 0.95em;
    }
    
    .fit-icon { font-size: 1.3em; margin-right: 0.5rem; }
    
    .metric-card {
        background: linear-gradient(135deg, #f5f7fa 0%, #ffffff 100%);
        border-left: 4px solid #0052CC;
        padding: 1.5rem;
        border-radius: 8px;
        box-shadow: 0 2px 8px rgba(0,0,0,0.05);
    }
    
    .metric-card-success {
        border-left-color: #28a745;
    }
    
    .metric-card-warning {
        border-left-color: #ffc107;
    }
    
    .risk-card {
        background: #fff3cd;
        border-left: 4px solid #ffc107;
        padding: 1rem;
        border-radius: 6px;
        margin: 0.5rem 0;
    }
    
    .flag-card {
        background: #f8d7da;
        border-left: 4px solid #dc3545;
        padding: 1rem;
        border-radius: 6px;
        margin: 0.5rem 0;
    }
    
    .section-header {
        font-size: 1.3em;
        font-weight: 600;
        color: #0052CC;
        margin-top: 1.5rem;
        margin-bottom: 1rem;
        border-bottom: 2px solid #e0e0e0;
        padding-bottom: 0.5rem;
    }
    
    .extracted-facts {
        background: #f9fafb;
        border-radius: 8px;
        padding: 1.5rem;
        margin-bottom: 1.5rem;
        border: 1px solid #e5e7eb;
    }
    
    .fact-row {
        display: grid;
        grid-template-columns: 1fr 1fr;
        gap: 1.5rem;
        margin-bottom: 1rem;
    }
    
    .fact-item {
        background: white;
        padding: 1rem;
        border-radius: 6px;
        border-left: 3px solid #3b82f6;
    }
    
    .fact-label {
        font-size: 0.85em;
        color: #6b7280;
        font-weight: 600;
        margin-bottom: 0.5rem;
    }
    
    .fact-value {
        color: #1f2937;
        line-height: 1.5;
    }
    
    .score-pill-good {
        display: inline-block;
        background: #d4edda;
        color: #155724;
        padding: 0.5rem 1rem;
        border-radius: 20px;
        font-weight: 600;
        margin: 0.5rem 0.25rem 0.5rem 0;
    }
    
    .quick-action-btn {
        display: inline-block;
        background: #f0f7ff;
        border: 1px solid #0052CC;
        color: #0052CC;
        padding: 0.6rem 1rem;
        border-radius: 6px;
        font-size: 0.85em;
        font-weight: 600;
        cursor: pointer;
        margin: 0.3rem;
    }
    
    .quick-action-btn:hover {
        background: #0052CC;
        color: white;
    }
    
    .checklist-item {
        display: flex;
        align-items: center;
        padding: 0.7rem 0;
        border-bottom: 1px solid #e5e7eb;
        font-size: 0.95em;
    }
    
    .checklist-item:last-child {
        border-bottom: none;
    }
    
    .status-badge {
        display: inline-block;
        padding: 0.35rem 0.75rem;
        border-radius: 20px;
        font-size: 0.75em;
        font-weight: 700;
        text-transform: uppercase;
        letter-spacing: 0.5px;
    }
    
    .status-new { background: #dbeafe; color: #1e40af; }
    .status-analyzed { background: #d4edda; color: #155724; }
    .status-needs-info { background: #fff3cd; color: #856404; }
    .status-ready { background: #d4edda; color: #155724; }
    
    .ic-memo {
        background: #f9fafb;
        border: 1px solid #e5e7eb;
        border-radius: 12px;
        padding: 2rem;
        font-family: 'Georgia', serif;
        line-height: 1.8;
        color: #1f2937;
    }
    
    .ic-memo h2 {
        border-bottom: 2px solid #0052CC;
        padding-bottom: 0.5rem;
        margin-top: 1.5rem;
    }
    
    .ic-memo h2:first-child {
        margin-top: 0;
    }
    
    </style>
""", unsafe_allow_html=True)

# ----------------------------
# Helper Functions
# ----------------------------

def mock_scorecard(deal: dict, investor_prefs: dict = None) -> dict:
    """Generate mock scorecard with reasoning."""
    base = 0.62
    reasoning = []
    
    if deal["stage"] in ["Seed", "Pre-Seed"]:
        base -= 0.05
        reasoning.append("Adjustment: Early stage companies carry execution risk")
    if deal["growth_rate_pct"] >= 15:
        base += 0.08
        reasoning.append(f"Boost: Strong growth rate ({deal['growth_rate_pct']}% is solid)")
    if deal["runway_months"] < 9:
        base -= 0.07
        reasoning.append("Concern: Short runway (<9 months) increases execution pressure")
    if deal["arr_usd"] >= 1_000_000:
        base += 0.06
        reasoning.append("Positive: Significant ARR indicates market traction")
    
    if investor_prefs:
        if investor_prefs.get("preferred_stage") == deal["stage"]:
            base += 0.05
            reasoning.append(f"Match: This aligns with your preferred stage ({deal['stage']})")
        if investor_prefs.get("preferred_sector") and investor_prefs["preferred_sector"].lower() in deal["sector"].lower():
            base += 0.05
            reasoning.append(f"Sector match: This deal is in {deal['sector']}, which matches your focus")
        if investor_prefs.get("min_arr") and deal["arr_usd"] >= investor_prefs["min_arr"]:
            base += 0.03
            reasoning.append(f"Revenue threshold: Exceeds your minimum ARR target")

    prob = max(0.05, min(0.95, base))
    drivers_pos = [
        "Strong growth rate (reported)",
        "Clear raise amount and use of funds",
        "Evidence of market pull / traction",
    ]
    drivers_neg = [
        "Short runway raises execution risk",
        "Missing retention / CAC details",
        "Team background unclear in deck",
    ]
    flags = ["Retention not provided", "CAC not provided"] if deal["arr_usd"] > 0 else ["No revenue metric provided"]

    return {
        "prob_next_round": prob,
        "confidence": 0.72,
        "drivers_pos": drivers_pos,
        "drivers_neg": drivers_neg,
        "risk_flags": flags,
        "reasoning": reasoning,
    }

def compute_thesis_fit(deal: dict, investor_prefs: dict) -> dict:
    """Compute thesis fit score based on investor preferences."""
    if not investor_prefs:
        return None
    
    score = 50
    matches = {"stage": None, "sector": None, "arr": None, "red_flags": []}
    
    # Stage match
    if investor_prefs.get("preferred_stage") and investor_prefs["preferred_stage"] == deal["stage"]:
        score += 20
        matches["stage"] = "✅"
    elif investor_prefs.get("preferred_stage"):
        matches["stage"] = "⚠️"
    else:
        matches["stage"] = "○"
    
    # Sector match
    if investor_prefs.get("preferred_sector") and investor_prefs["preferred_sector"].lower() in deal["sector"].lower():
        score += 20
        matches["sector"] = "✅"
    elif investor_prefs.get("preferred_sector"):
        matches["sector"] = "⚠️"
    else:
        matches["sector"] = "○"
    
    # ARR threshold
    if investor_prefs.get("min_arr") and deal["arr_usd"] >= investor_prefs["min_arr"]:
        score += 10
        matches["arr"] = "✅"
    elif investor_prefs.get("min_arr"):
        matches["arr"] = "❌"
    else:
        matches["arr"] = "○"
    
    # Red flags check
    if deal["runway_months"] < 6:
        matches["red_flags"].append("Very short runway (<6 months)")
        score -= 15
    
    score = max(0, min(100, score))
    
    return {
        "score": score,
        "stage_match": matches["stage"],
        "sector_match": matches["sector"],
        "arr_match": matches["arr"],
        "red_flags": matches["red_flags"],
    }

def generate_recommendation(prob: float, thesis_fit: dict, investor_prefs: dict) -> dict:
    """Generate VC-style recommendation based on prob + thesis fit."""
    combined_signal = prob * 0.6 + (thesis_fit["score"] / 100) * 0.4 if thesis_fit else prob
    
    if combined_signal >= 0.70:
        rec = "INVEST"
        confidence = "High" if combined_signal >= 0.80 else "Medium"
    elif combined_signal >= 0.50:
        rec = "WATCHLIST"
        confidence = "Medium" if combined_signal >= 0.60 else "Low"
    else:
        rec = "PASS"
        confidence = "Medium" if combined_signal >= 0.40 else "Low"
    
    return {
        "recommendation": rec,
        "confidence": confidence,
        "model_signal": prob,
        "combined_signal": combined_signal,
    }

def generate_extracted_highlights(docs_text: str, client: OpenAI) -> dict:
    """Extract key business facts from documents."""
    if not docs_text or len(docs_text.strip()) < 100:
        return None
    
    system = "Extract key business facts from the pitch deck. Return ONLY valid JSON with these fields: business_model, icp, pricing, gtm_strategy, traction_metrics, moat, competition, team_highlights. If not found, use null."
    user = f"""
Extract business highlights from:

{docs_text[:8000]}

Return JSON with: business_model, icp, pricing, gtm_strategy, traction_metrics, moat, competition, team_highlights
"""
    
    try:
        resp = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": user},
            ],
            response_format={"type": "json_object"},
            temperature=0.5,
        )
        return json.loads(resp.choices[0].message.content.strip())
    except:
        return None

def generate_founder_questions(deal: dict, result: dict, client: OpenAI) -> str:
    """Generate 8-12 specific diligence questions for founder."""
    system = "You are a VC conducting diligence. Generate 8-12 specific, insightful questions for the founder based on their deal profile."
    user = f"""
Company: {deal['company']}
Stage: {deal['stage']}
Sector: {deal['sector']}
ARR: ${deal['arr_usd']:,}
Growth: {deal['growth_rate_pct']}%
Runway: {deal['runway_months']} months

Model risks: {result['drivers_neg']}
Missing info: {result['risk_flags']}

Generate specific founder questions (8-12, numbered):
"""
    
    try:
        resp = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": user},
            ],
            temperature=0.7,
        )
        return resp.choices[0].message.content.strip()
    except:
        return "Could not generate questions."

def generate_ic_memo(deal: dict, result: dict, thesis_fit: dict, recommendation: dict, client: OpenAI) -> str:
    """Generate formal IC memo."""
    system = "You are a VC writing an investment committee memo. Be concise, professional, and data-driven."
    user = f"""
Generate an IC memo for this deal:

Company: {deal['company']}
Stage: {deal['stage']}
Sector: {deal['sector']}
Raise: ${deal['raise_amount_usd']:,}
ARR: ${deal['arr_usd']:,}
Growth: {deal['growth_rate_pct']}%

Model Probability: {result['prob_next_round']:.0%}
Recommendation: {recommendation['recommendation']}
Thesis Fit: {thesis_fit['score']}/100

Drivers (+): {', '.join(result['drivers_pos'])}
Drivers (-): {', '.join(result['drivers_neg'])}

Write a structured IC memo with sections:
# [Company Name] – Investment Memo
## Overview
## Stage & Market
## Traction
## Team & Execution
## Risks & Concerns
## Investment Thesis
## Recommendation

Keep it to 2-3 sentences per section.
"""
    
    try:
        resp = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": user},
            ],
            temperature=0.6,
        )
        return resp.choices[0].message.content.strip()
    except:
        return "Could not generate memo."

def apply_conviction(prob: float, conviction: float) -> float:
    """Adjust probability by conviction slider."""
    return max(0.0, min(1.0, prob + 0.08 * conviction))

def pdf_to_text(file) -> str:
    import fitz
    data = file.read()
    doc = fitz.open(stream=data, filetype="pdf")
    parts = []
    for page in doc:
        parts.append(page.get_text("text"))
    return "\n".join(parts)

def txt_to_text(file) -> str:
    return file.read().decode("utf-8", errors="ignore")

def tabular_to_text(file) -> str:
    import pandas as pd
    name = (file.name or "").lower()
    if name.endswith(".csv"):
        df = pd.read_csv(file)
    else:
        df = pd.read_excel(file)
    head = df.head(50)
    return head.to_csv(index=False)

def pptx_to_text(file) -> str:
    from pptx import Presentation
    import io
    data = file.read()
    prs = Presentation(io.BytesIO(data))
    parts = []
    for i, slide in enumerate(prs.slides, start=1):
        parts.append(f"\n--- SLIDE {i} ---")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text)
    return "\n".join(parts)

def stage_apply_extracted(ex: dict):
    """Stage extracted values for application on rerun."""
    st.session_state["_pending_apply"] = ex
    st.rerun()

# ----------------------------
# Session State Initialization
# ----------------------------
if "investor_prefs" not in st.session_state:
    st.session_state.investor_prefs = None

if "show_prefs_onboard" not in st.session_state:
    st.session_state.show_prefs_onboard = False

if "extracted_highlights" not in st.session_state:
    st.session_state.extracted_highlights = None

if "founder_questions" not in st.session_state:
    st.session_state.founder_questions = None

if "ic_memo" not in st.session_state:
    st.session_state.ic_memo = None

# ----------------------------
# Sidebar: Investor Profile
# ----------------------------
with st.sidebar:
    st.markdown("### ⚙️ Investor Profile")
    
    if not st.session_state.investor_prefs:
        st.info("👋 Welcome! Configure your preferences to personalize scoring.", icon="ℹ️")
        if st.button("→ Configure Preferences", key="open_prefs", use_container_width=True):
            st.session_state.show_prefs_onboard = True
            st.rerun()
    else:
        prefs = st.session_state.investor_prefs
        st.success("✅ Preferences configured", icon="✓")
        with st.expander("View / Edit Profile"):
            st.write(f"**Name:** {prefs.get('investor_name', 'Not set')}")
            st.write(f"**Fund Focus:** {prefs.get('fund_type', 'Not set')}")
            st.write(f"**Preferred Stage:** {prefs.get('preferred_stage', 'Not set')}")
            st.write(f"**Preferred Sector:** {prefs.get('preferred_sector', 'Not set')}")
            st.write(f"**Deal Types:** {prefs.get('deal_types', 'Not set')}")
            if st.button("Edit Preferences", use_container_width=True):
                st.session_state.show_prefs_onboard = True
                st.rerun()
            if st.button("Clear Preferences", use_container_width=True):
                st.session_state.investor_prefs = None
                st.rerun()
    
    st.divider()
    st.markdown("### 📊 About")
    st.caption("VCaaS extracts deal metrics, scores likelihood of next round, and answers questions about your deals.")

# ----------------------------
# Onboarding: Investor Preferences
# ----------------------------
if st.session_state.show_prefs_onboard:
    st.markdown("---")
    st.markdown("## 🎯 Investor Preferences & Ethos")
    st.markdown("Tell us about yourself (optional but recommended). This helps personalize deal scoring.")
    
    col1, col2 = st.columns(2)
    with col1:
        investor_name = st.text_input(
            "Your name / Fund name",
            value=st.session_state.investor_prefs.get("investor_name", "") if st.session_state.investor_prefs else "",
            placeholder="e.g., John Smith, Acme Ventures"
        )
        preferred_stage = st.selectbox(
            "Primary investment stage",
            ["Any", "Pre-Seed", "Seed", "Series A", "Series B+"],
            index=0,
            key="stage_select"
        )
    
    with col2:
        fund_type = st.selectbox(
            "Fund type / Structure",
            ["Any", "VC", "Angel", "Micro VC", "Corporate VC", "PE"],
            index=0,
            key="fund_select"
        )
        min_arr = st.number_input(
            "Minimum ARR for interest ($)",
            min_value=0,
            step=100000,
            value=st.session_state.investor_prefs.get("min_arr", 0) if st.session_state.investor_prefs else 0,
            help="Deals below this ARR won't match your profile"
        )
    
    preferred_sector = st.text_input(
        "Preferred sector(s)",
        value=st.session_state.investor_prefs.get("preferred_sector", "") if st.session_state.investor_prefs else "",
        placeholder="e.g., B2B SaaS, AI, Climate Tech, or comma-separated"
    )
    
    deal_types = st.multiselect(
        "Deal types you're strong in",
        ["Product-led growth", "Sales-led growth", "Marketplace", "Infrastructure", "Consumer", "Deep tech", "Enterprise"],
        default=st.session_state.investor_prefs.get("deal_types", []) if st.session_state.investor_prefs else []
    )
    
    conviction_style = st.radio(
        "How much detail do you like in analysis?",
        ["Concise summaries", "Balanced", "Deep dives with lots of detail"],
        index=1
    )
    
    st.markdown("**Investment ethos / red flags:**")
    ethos = st.text_area(
        "What do you look for? What are your hard passes?",
        value=st.session_state.investor_prefs.get("ethos", "") if st.session_state.investor_prefs else "",
        placeholder="e.g., Love capital-efficient teams, strong unit economics. Hard pass on consumer social. Must have experienced CTO.",
        height=100
    )
    
    col_save, col_cancel = st.columns(2)
    with col_save:
        if st.button("💾 Save Preferences", use_container_width=True, type="primary"):
            st.session_state.investor_prefs = {
                "investor_name": investor_name,
                "preferred_stage": preferred_stage if preferred_stage != "Any" else None,
                "fund_type": fund_type if fund_type != "Any" else None,
                "preferred_sector": preferred_sector,
                "min_arr": min_arr,
                "deal_types": deal_types,
                "conviction_style": conviction_style,
                "ethos": ethos,
            }
            st.session_state.show_prefs_onboard = False
            st.success("✅ Preferences saved!", icon="✓")
            st.rerun()
    
    with col_cancel:
        if st.button("Cancel", use_container_width=True):
            st.session_state.show_prefs_onboard = False
            st.rerun()
    
    st.markdown("---")

# ----------------------------
# App Header
# ----------------------------
st.markdown("<h1 class='main-title'>💼 VCaaS – Deal Dashboard</h1>", unsafe_allow_html=True)
st.caption("📤 Upload pitch decks → 📊 Get instant scoring → 💬 Chat with your copilot")

# ----------------------------
# API Key Setup
# ----------------------------
openai_api_key = st.secrets.get("OPENAI_API_KEY", "")
if not openai_api_key:
    openai_api_key = st.text_input("OpenAI API Key", type="password")

if not openai_api_key:
    st.info("Add your OpenAI API key to continue.", icon="🗝️")
    st.stop()

client = OpenAI(api_key=openai_api_key)

# ----------------------------
# Layout: 2x2 Grid (Docs | Fields / Scorecard | Copilot)
# ----------------------------
col_docs, col_fields = st.columns([1.0, 0.9], gap="large")

# ============================================
# 1A) DOCUMENTS UPLOAD (left, top)
# ============================================
if "extracted" not in st.session_state:
    st.session_state.extracted = {}

with col_docs:
    st.subheader("📄 Pitch Deck & Docs")
    
    deck_file = st.file_uploader("Pitch deck (PDF / PPTX)", type=["pdf", "pptx"], key="deck")
    extra_docs = st.file_uploader(
        "Additional materials (optional)",
        type=["pdf", "pptx", "txt", "md"],
        accept_multiple_files=True,
        key="extra",
    )
    financials_file = st.file_uploader(
        "Financials (optional) — CSV / XLSX",
        type=["csv", "xlsx"],
        key="fin",
    )
    notes = st.text_area(
        "Additional context (optional)",
        height=80,
        placeholder="Paste extra notes: founder email, meeting notes, traction metrics, objections, diligence notes, etc.",
        key="notes",
    )
    st.caption(f"📝 {len(st.session_state.get('notes','').strip()):,} characters")

    # Build combined docs text
    docs_parts = []

    if deck_file is not None:
        try:
            name = (deck_file.name or "").lower()
            if name.endswith(".pdf"):
                text = pdf_to_text(deck_file)
            elif name.endswith(".pptx"):
                text = pptx_to_text(deck_file)
            else:
                text = ""
            docs_parts.append("===== PITCH DECK =====\n" + text)
            st.success("✓ Pitch deck loaded")
        except Exception as e:
            st.error(f"Could not read pitch deck: {e}")

    if extra_docs:
        for f in extra_docs:
            try:
                name = (f.name or "").lower()
                if name.endswith(".pdf"):
                    text = pdf_to_text(f)
                elif name.endswith(".pptx"):
                    text = pptx_to_text(f)
                else:
                    text = txt_to_text(f)
                docs_parts.append(f"===== EXTRA DOC: {f.name} =====\n{text}")
            except Exception as e:
                st.error(f"Could not read {f.name}: {e}")

    if financials_file is not None:
        try:
            tabular_text = tabular_to_text(financials_file)
            docs_parts.append(
                f"===== FINANCIALS PREVIEW: {financials_file.name} (top rows) =====\n{tabular_text}"
            )
            st.success("✓ Financials loaded")
        except Exception as e:
            st.error(f"Could not read financials: {e}")

    if st.session_state.get("notes", "").strip():
        docs_parts.append("===== ADDITIONAL CONTEXT (TEXT) =====\n" + st.session_state["notes"].strip())

    combined_docs_text = "\n\n".join(docs_parts)

    st.divider()

    col_extract, col_preview = st.columns([1, 1])
    with col_extract:
        extract = st.button("🤖 Extract fields", use_container_width=True, type="primary")

    with col_preview:
        with st.expander("🔍 Preview"):
            st.caption(f"Characters: {len(combined_docs_text):,}")
            st.text(combined_docs_text[:3000] if combined_docs_text else "No text yet.")

    if extract:
        if not combined_docs_text.strip():
            st.warning("Upload a deck or add context first.")
        else:
            with st.spinner("Extracting deal fields..."):
                system = "Extract deal fields from the provided documents. Output ONLY valid JSON."
                user = f"""
Return JSON with these keys:
{list(EXTRACT_SCHEMA.keys())}

Rules:
- Use numbers for numeric fields (no $ or commas).
- If unknown, use null.
- stage should be one of: Pre-Seed, Seed, Series A, Series B+
- Keep it conservative: don't guess.

DOCUMENTS:
{combined_docs_text[:12000]}
"""
                resp = client.chat.completions.create(
                    model="gpt-4o-mini",
                    messages=[
                        {"role": "system", "content": system},
                        {"role": "user", "content": user},
                    ],
                    response_format={"type": "json_object"},
                )

                raw = resp.choices[0].message.content.strip()
                try:
                    extracted = json.loads(raw)
                    st.session_state.extracted = extracted
                    st.success("✓ Extraction complete!")
                    with st.expander("View extracted JSON"):
                        st.json(extracted)
                except Exception:
                    st.error("Extraction failed (non-JSON returned).")
                    st.code(raw)

    if st.session_state.get("extracted"):
        if st.button("→ Apply to fields →", use_container_width=True):
            stage_apply_extracted(st.session_state.extracted)

# ============================================
# 1B) DEAL FIELDS (right, top)
# ============================================
with col_fields:
    st.subheader("📋 Deal Info")
    
    if "_pending_apply" in st.session_state:
        ex = st.session_state.pop("_pending_apply") or {}
        for k in ["company","stage","sector","raise_amount_usd","arr_usd","growth_rate_pct","runway_months","notes"]:
            v = ex.get(k, None)
            if v is not None:
                st.session_state[k] = v
    
    company = st.text_input("Company name", value="Acme AI", key="company")
    
    col1, col2 = st.columns(2)
    with col1:
        stage = st.selectbox("Stage", ["Pre-Seed", "Seed", "Series A", "Series B+"], key="stage")
    with col2:
        sector = st.text_input("Sector", value="B2B SaaS", key="sector")
    
    col1, col2 = st.columns(2)
    with col1:
        raise_amount = st.number_input("Raise amount ($)", min_value=0, step=250000, value=3000000, key="raise_amount_usd")
    with col2:
        arr = st.number_input("ARR ($)", min_value=0, step=100000, value=800000, key="arr_usd")
    
    col1, col2 = st.columns(2)
    with col1:
        growth = st.number_input("Growth rate %", min_value=0, max_value=500, step=1, value=18, key="growth_rate_pct")
    with col2:
        runway = st.number_input("Runway (months)", min_value=0, max_value=60, step=1, value=10, key="runway_months")

    st.divider()
    run = st.button("▶ Run VCaaS analysis", type="primary", use_container_width=True)

st.divider()

# ============================================
# 2) SCORECARD & COPILOT (Row 2)
# ============================================
col_score, col_chat = st.columns([1.0, 1.2], gap="large")

# ============================================
# 2A) SCORECARD (left, bottom) - PREMIUM VERSION
# ============================================
with col_score:
    st.subheader("📊 Scorecard & Analysis")

    if "last_result" not in st.session_state:
        st.session_state.last_result = None
        st.session_state.last_deal = None

    if run:
        deal = {
            "company": company,
            "stage": stage,
            "sector": sector,
            "raise_amount_usd": raise_amount,
            "arr_usd": arr,
            "growth_rate_pct": growth,
            "runway_months": runway,
            "docs_text": combined_docs_text,
            "notes": notes,
            "has_deck": deck_file is not None,
            "extra_doc_count": len(extra_docs) if extra_docs else 0,
            "has_financials": financials_file is not None,
        }

        st.session_state.last_deal = deal
        st.session_state.last_result = mock_scorecard(deal, st.session_state.investor_prefs)
        st.session_state.extracted_highlights = None
        st.session_state.founder_questions = None
        st.session_state.ic_memo = None

    if not st.session_state.last_result:
        st.info("👈 Fill in the form on the left and click **Run VCaaS analysis** to begin.")
    else:
        deal = st.session_state.last_deal
        result = st.session_state.last_result

        # ---- DEAL SNAPSHOT (Premium) ----
        st.markdown("""<div class='deal-snapshot'>""", unsafe_allow_html=True)
        
        snapshot_company = deal['company']
        snapshot_tagline = "Enterprise SaaS platform"  # Could extract from docs
        
        st.markdown(f"<div class='snapshot-header'>{snapshot_company}</div>", unsafe_allow_html=True)
        st.markdown(f"<div class='snapshot-tagline'>{snapshot_tagline}</div>", unsafe_allow_html=True)
        
        snap_col1, snap_col2, snap_col3, snap_col4 = st.columns(4)
        with snap_col1:
            st.markdown(f"""
            <div class='snapshot-item'>
                <div class='snapshot-label'>Stage</div>
                <div class='snapshot-value'>{deal['stage']}</div>
            </div>
            """, unsafe_allow_html=True)
        with snap_col2:
            st.markdown(f"""
            <div class='snapshot-item'>
                <div class='snapshot-label'>Sector</div>
                <div class='snapshot-value'>{deal['sector']}</div>
            </div>
            """, unsafe_allow_html=True)
        with snap_col3:
            st.markdown(f"""
            <div class='snapshot-item'>
                <div class='snapshot-label'>Raise</div>
                <div class='snapshot-value'>${deal['raise_amount_usd']/1e6:.1f}M</div>
            </div>
            """, unsafe_allow_html=True)
        with snap_col4:
            st.markdown(f"""
            <div class='snapshot-item'>
                <div class='snapshot-label'>ARR</div>
                <div class='snapshot-value'>${deal['arr_usd']/1e6:.1f}M</div>
            </div>
            """, unsafe_allow_html=True)
        
        snap_col1, snap_col2, snap_col3, snap_col4 = st.columns(4)
        with snap_col1:
            st.markdown(f"""
            <div class='snapshot-item'>
                <div class='snapshot-label'>Growth</div>
                <div class='snapshot-value'>{deal['growth_rate_pct']}%</div>
            </div>
            """, unsafe_allow_html=True)
        with snap_col2:
            st.markdown(f"""
            <div class='snapshot-item'>
                <div class='snapshot-label'>Runway</div>
                <div class='snapshot-value'>{deal['runway_months']} mo</div>
            </div>
            """, unsafe_allow_html=True)
        with snap_col3:
            docs_len = len(deal.get('docs_text', ''))
            st.markdown(f"""
            <div class='snapshot-item'>
                <div class='snapshot-label'>Docs</div>
                <div class='snapshot-value'>{docs_len/1000:.0f}K chars</div>
            </div>
            """, unsafe_allow_html=True)
        with snap_col4:
            has_all = "✓" if (deal['has_deck'] and deal['has_financials']) else "⚠️"
            st.markdown(f"""
            <div class='snapshot-item'>
                <div class='snapshot-label'>Materials</div>
                <div class='snapshot-value'>{has_all}</div>
            </div>
            """, unsafe_allow_html=True)
        
        st.markdown("""</div>""", unsafe_allow_html=True)

        # ---- RECOMMENDATION CARD (VC DECISION) ----
        thesis_fit = compute_thesis_fit(deal, st.session_state.investor_prefs) if st.session_state.investor_prefs else None
        recommendation = generate_recommendation(result["prob_next_round"], thesis_fit, st.session_state.investor_prefs)

        rec_color_map = {
            "INVEST": "recommendation-badge-invest",
            "WATCHLIST": "recommendation-badge-watchlist",
            "PASS": "recommendation-badge-pass",
        }
        conf_color_map = {
            "High": "confidence-high",
            "Medium": "confidence-med",
            "Low": "confidence-low",
        }

        st.markdown(f"""<div class='recommendation-card'>
        <div class='recommendation-label'>Investment Recommendation</div>
        <div class='{rec_color_map[recommendation["recommendation"]]}'>{recommendation["recommendation"]}</div>
        <div style='margin-top: 1rem; font-size: 1em;'>
            Confidence: <span class='{conf_color_map[recommendation["confidence"]]}'>{recommendation["confidence"]}</span><br>
            Model Signal: <strong>{recommendation['model_signal']*100:.0f}%</strong>
        </div>
        </div>""", unsafe_allow_html=True)

        # ---- THESIS FIT SCORE ----
        if thesis_fit:
            fit_color = "fit-score-good" if thesis_fit["score"] >= 70 else "fit-score-med" if thesis_fit["score"] >= 50 else "fit-score-bad"
            st.markdown(f"""<div class='thesis-fit-score'>
            <div style='text-align: center; font-size: 0.9em; color: #6b7280; text-transform: uppercase; letter-spacing: 1px; margin-bottom: 0.5rem;'>Thesis Fit Score</div>
            <div class='fit-gauge {fit_color}'>{thesis_fit["score"]}/100</div>
            <div class='fit-item'><span class='fit-icon'>{thesis_fit["stage_match"]}</span> Stage match</div>
            <div class='fit-item'><span class='fit-icon'>{thesis_fit["sector_match"]}</span> Sector match</div>
            <div class='fit-item'><span class='fit-icon'>{thesis_fit["arr_match"]}</span> Revenue threshold</div>
            """, unsafe_allow_html=True)
            if thesis_fit["red_flags"]:
                st.markdown("**🚩 Red flags:**", unsafe_allow_html=True)
                for flag in thesis_fit["red_flags"]:
                    st.markdown(f"• {flag}", unsafe_allow_html=True)
            st.markdown("</div>", unsafe_allow_html=True)

        st.divider()

        # ---- CONVICTION SLIDER ----
        conviction = st.slider(
            "🎯 Conviction / Vibe adjustment",
            min_value=-1.0, max_value=1.0, value=0.0, step=0.1,
            help="Adjust score based on your gut feel about the team or market"
        )
        adjusted = apply_conviction(result["prob_next_round"], conviction)

        col_m1, col_m2 = st.columns(2)
        with col_m1:
            st.markdown(f"""
            <div class='metric-card metric-card-success'>
                <div style='font-size: 0.9em; color: #666; margin-bottom: 0.5rem;'>📊 Model Probability</div>
                <div style='font-size: 2.2em; font-weight: 700; color: #0052CC;'>{result['prob_next_round']*100:.0f}%</div>
            </div>
            """, unsafe_allow_html=True)
        
        with col_m2:
            score_color = "#28a745" if adjusted >= 0.65 else "#ffc107" if adjusted >= 0.5 else "#dc3545"
            st.markdown(f"""
            <div class='metric-card' style='border-left-color: {score_color};'>
                <div style='font-size: 0.9em; color: #666; margin-bottom: 0.5rem;'>🎯 Adjusted</div>
                <div style='font-size: 2.2em; font-weight: 700; color: {score_color};'>{adjusted*100:.0f}%</div>
            </div>
            """, unsafe_allow_html=True)
        
        st.progress(adjusted)

        st.divider()

        # ---- MODEL REASONING (Explainability) ----
        st.markdown(f"<div class='section-header'>🧠 Model Reasoning</div>", unsafe_allow_html=True)
        with st.expander("💭 Show me the model's thinking", expanded=True):
            st.write("**How the score was calculated:**")
            if result.get("reasoning"):
                for reason in result["reasoning"]:
                    st.write(f"→ {reason}")
            st.divider()
            st.write("**Confidence level:** This model is " + ("very confident" if result['confidence'] > 0.8 else "moderately confident" if result['confidence'] > 0.65 else "less confident") + " in this assessment.")

        # ---- EXTRACTED HIGHLIGHTS ----
        if not st.session_state.extracted_highlights and combined_docs_text:
            if st.button("✨ Extract key business facts", use_container_width=True):
                with st.spinner("Analyzing documents..."):
                    st.session_state.extracted_highlights = generate_extracted_highlights(combined_docs_text, client)

        if st.session_state.extracted_highlights:
            st.markdown(f"<div class='section-header'>✨ Key Extracted Facts</div>", unsafe_allow_html=True)
            highlights = st.session_state.extracted_highlights
            st.markdown("<div class='extracted-facts'>", unsafe_allow_html=True)
            
            for key in ["business_model", "icp", "pricing", "gtm_strategy"]:
                if highlights.get(key):
                    st.markdown(f"**{key.replace('_', ' ').title()}:**\n{highlights[key]}")
            
            st.markdown("</div>", unsafe_allow_html=True)

        st.divider()

        # ---- POSITIVE & NEGATIVE DRIVERS ----
        st.markdown(f"<div class='section-header'>✅ Positive Drivers</div>", unsafe_allow_html=True)
        for d in result["drivers_pos"]:
            st.markdown(f"<span class='score-pill-good'>✅ {d}</span>", unsafe_allow_html=True)

        st.markdown(f"<div class='section-header'>⚠️ Risk Drivers</div>", unsafe_allow_html=True)
        for d in result["drivers_neg"]:
            st.markdown(f"<div class='risk-card'>⚠️ {d}</div>", unsafe_allow_html=True)

        st.markdown(f"<div class='section-header'>🚩 Missing Info & Flags</div>", unsafe_allow_html=True)
        for f in result["risk_flags"]:
            st.markdown(f"<div class='flag-card'>🚩 {f}</div>", unsafe_allow_html=True)

        st.divider()

        # ---- FOUNDER QUESTIONS BUTTON ----
        col_founder, col_memo = st.columns(2)
        with col_founder:
            if not st.session_state.founder_questions:
                if st.button("❓ Generate founder questions", use_container_width=True):
                    with st.spinner("Generating diligence questions..."):
                        st.session_state.founder_questions = generate_founder_questions(deal, result, client)
            else:
                if st.button("🔄 Regenerate questions", use_container_width=True):
                    with st.spinner("Generating diligence questions..."):
                        st.session_state.founder_questions = generate_founder_questions(deal, result, client)

        # ---- IC MEMO BUTTON ----
        with col_memo:
            if not st.session_state.ic_memo:
                if st.button("📋 Generate IC memo", use_container_width=True):
                    with st.spinner("Drafting IC memo..."):
                        st.session_state.ic_memo = generate_ic_memo(deal, result, thesis_fit or {}, recommendation, client)
            else:
                if st.button("🔄 Regenerate memo", use_container_width=True):
                    with st.spinner("Drafting IC memo..."):
                        st.session_state.ic_memo = generate_ic_memo(deal, result, thesis_fit or {}, recommendation, client)

        # ---- DISPLAY FOUNDER QUESTIONS ----
        if st.session_state.founder_questions:
            with st.expander("❓ Founder Diligence Questions", expanded=False):
                st.markdown(st.session_state.founder_questions)

        # ---- DISPLAY IC MEMO ----
        if st.session_state.ic_memo:
            with st.expander("📋 Investment Committee Memo", expanded=False):
                st.markdown(f"<div class='ic-memo'>\n{st.session_state.ic_memo}\n</div>", unsafe_allow_html=True)

# ============================================
# 2B) COPILOT (right, bottom)
# ============================================
with col_chat:
    st.subheader("3) 💬 VCaaS Copilot")

    if "chat" not in st.session_state:
        st.session_state.chat = []

    # Quick action buttons
    st.markdown("**Quick Actions:**")
    col1, col2, col3 = st.columns(3)
    
    quick_actions = [
        ("📋 Summarize deck", "Summarize this pitch deck in 8 key bullet points. Focus on business model, market, traction, and team."),
        ("⚠️ List risks", "What are the top 5-8 diligence risks for this deal? Be specific."),
        ("📈 Traction metrics", "Extract and summarize all traction metrics from the pitch deck."),
    ]
    
    col_idx = 0
    for i, (label, prompt_template) in enumerate(quick_actions):
        col = [col1, col2, col3][i % 3]
        with col:
            if st.button(label, use_container_width=True, key=f"qa_{i}"):
                st.session_state.chat.append({"role": "user", "content": prompt_template})
    
    col1, col2, col3 = st.columns(3)
    
    quick_actions_2 = [
        ("📧 Founder follow-up", "Draft a professional follow-up email to the founder outlining next steps and key questions from this analysis."),
        ("📊 IC Summary", "Write a 3-paragraph investment committee summary for this deal."),
        ("🔍 Next steps", "What should we do next to advance this deal? Suggest specific diligence items and timeline."),
    ]
    
    for i, (label, prompt_template) in enumerate(quick_actions_2):
        col = [col1, col2, col3][i % 3]
        with col:
            if st.button(label, use_container_width=True, key=f"qa_{i+3}"):
                st.session_state.chat.append({"role": "user", "content": prompt_template})

    st.divider()

    # Chat history
    for msg in st.session_state.chat:
        with st.chat_message(msg["role"]):
            st.markdown(msg["content"])

    # Chat input
    prompt = st.chat_input("Ask about the deck, the deal, or the scorecard…")

    if prompt:
        if not st.session_state.last_deal or not st.session_state.last_result:
            st.warning("Run the analysis first so I have deal context + a scorecard.")
        else:
            deal = st.session_state.last_deal
            result = st.session_state.last_result

            st.session_state.chat.append({"role": "user", "content": prompt})
            with st.chat_message("user"):
                st.markdown(prompt)

            # Copilot system prompt with investor context
            investor_context = ""
            if st.session_state.investor_prefs:
                prefs = st.session_state.investor_prefs
                investor_context = f"""

INVESTOR PROFILE:
- Investor: {prefs.get('investor_name', 'Not specified')}
- Preferred stage: {prefs.get('preferred_stage', 'Any')}
- Preferred sector: {prefs.get('preferred_sector', 'Any')}
- Fund type: {prefs.get('fund_type', 'Any')}
- Investment ethos: {prefs.get('ethos', 'Not specified')}
- Deal types of interest: {', '.join(prefs.get('deal_types', ['Any'])) or 'Any'}
"""

            system = (
                "You are a VC/PE diligence copilot with deep deal analysis expertise. "
                "Be conversational, insightful, and specific. "
                "Use ONLY the provided deal fields, pitch deck text, and model outputs. "
                "Reference the model's reasoning and scoring when relevant. "
                f"Help evaluate deals through this investor's lens.{investor_context if investor_context else ''} "
                "If something isn't provided, say it's missing and suggest what to ask for or dig deeper on."
            )
            
            context = f"""
DEAL FIELDS (user-entered):
- Company: {deal['company']}
- Stage: {deal['stage']}
- Sector: {deal['sector']}
- Raise amount: ${deal['raise_amount_usd']:,}
- ARR: ${deal['arr_usd']:,}
- Growth rate %: {deal['growth_rate_pct']}
- Runway months: {deal['runway_months']}

MODEL OUTPUTS & REASONING:
- Probability of next round: {result['prob_next_round']:.2%}
- Confidence: {result['confidence']:.0%}
- Model reasoning:
"""
            
            for reason in result.get("reasoning", []):
                context += f"\n  • {reason}"
            
            context += f"""

KEY DRIVERS (Positive):
{chr(10).join('  • ' + d for d in result['drivers_pos'])}

KEY DRIVERS (Negative):
{chr(10).join('  • ' + d for d in result['drivers_neg'])}

MISSING INFO / FLAGS:
{chr(10).join('  • ' + f for f in result['risk_flags'])}

DOCUMENTS (deck + extras + financials preview + notes):
{deal.get('docs_text','')[:12000]}
"""

            messages = [
                {"role": "system", "content": system},
                {"role": "user", "content": context},
                {"role": "user", "content": prompt},
            ]

            with st.chat_message("assistant"):
                stream = client.chat.completions.create(
                    model="gpt-4o-mini",
                    messages=messages,
                    stream=True,
                )
                response = st.write_stream(stream)

            st.session_state.chat.append({"role": "assistant", "content": response})
